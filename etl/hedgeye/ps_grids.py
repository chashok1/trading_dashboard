"""
ps_grids.py — P5 and PTn grid writer for Portfolio Solutions emails.

Called automatically during PS dispatch (generate_from_email), and also
usable standalone to generate files for a specific date via IMAP lookup.

Usage:
    python -m etl.hedgeye.ps_grids 2026-06-26
"""
from __future__ import annotations

import email as email_lib
import imaplib
import logging
import os
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

log = logging.getLogger("hedgeye.ps_grids")

P5_DIR  = Path(r"C:\Ashok\Investing\Stocks\P5\Archive")
PTN_DIR = Path(r"C:\Ashok\Investing\Stocks\PTn\Archive")

_IMAP_MONTHS = ["Jan","Feb","Mar","Apr","May","Jun",
                 "Jul","Aug","Sep","Oct","Nov","Dec"]


def _imap_date(d: date) -> str:
    return f"{d.day:02d}-{_IMAP_MONTHS[d.month-1]}-{d.year}"


def _connect():
    from dotenv import load_dotenv
    load_dotenv()
    from etl.hedgeye.config import load
    cfg = load()
    mail = imaplib.IMAP4_SSL(cfg.imap_host)
    pwd = os.environ.get("HEDGEYE_IMAP_PASSWORD") or cfg.imap_password
    mail.login(cfg.imap_user, pwd)
    mail.select(cfg.mailbox)
    return mail


def _fetch_ps_html(mail, target_date: date) -> tuple[str, str] | None:
    since  = _imap_date(target_date)
    before = _imap_date(target_date + timedelta(days=3))
    _, ids = mail.search(None,
        f'SINCE {since} BEFORE {before} SUBJECT "PORTFOLIO SOLUTIONS"')
    id_list = ids[0].split()
    if not id_list:
        return None
    date_subj = f"{target_date.month}/{target_date.day}/{target_date.year}"
    for eid in id_list:
        _, data = mail.fetch(eid, "(RFC822)")
        msg = email_lib.message_from_bytes(data[0][1])
        subj = msg.get("Subject", "")
        if date_subj not in subj:
            continue
        html = ""
        for part in msg.walk():
            if part.get_content_type() == "text/html":
                html = part.get_payload(decode=True).decode("utf-8", errors="replace")
                break
        if html:
            return subj, html
    return None


def _parse_tables(html: str) -> list[list[list[str]]]:
    from etl.hedgeye.parsers import _TableExtractor
    ext = _TableExtractor()
    ext.feed(html)
    return ext.tables


def _is_p5(header: list[str]) -> bool:
    """P5: first col Ticker + subsequent cols are date labels like '26-JUN'."""
    return (
        bool(header)
        and header[0].strip().lower() == "ticker"
        and any(re.match(r"\d+-[A-Z]{3}$", h.strip()) for h in header[1:])
    )


def _is_ptn(header: list[str]) -> bool:
    """PTn: Ticker + Today + '1 Day Ago' + '1 Week Ago'."""
    low = [h.strip().lower() for h in header]
    return (
        "ticker" in low
        and "today" in low
        and "1 day ago" in low
        and "1 week ago" in low
    )


def _int(val: str) -> int | None:
    try:
        return int(val.strip())
    except (ValueError, AttributeError):
        return None


def _write_p5(tbl: list[list[str]], file_date: date, out: Path) -> int:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Sheet"
    ws.append(["Date", " TICKER", " TODAY", "ONEDAYAGO",
               "TWODAYAGO", "THREEDAYAGO", "FOURDAYAGO"])
    n = 0
    for row in tbl[1:]:
        if len(row) < 6:
            continue
        ticker = row[0].strip().upper()
        if not ticker:
            continue
        ws.append([
            datetime(file_date.year, file_date.month, file_date.day),
            ticker, _int(row[1]), _int(row[2]),
            _int(row[3]), _int(row[4]), _int(row[5]),
        ])
        n += 1
    wb.save(out)
    return n


def _write_ptn(tbl: list[list[str]], file_date: date, out: Path) -> int:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Sheet"
    ws.append(["Date", " TICKER", "TODAY", "1 DAY AGO",
               "1 WEEK AGO", "1 MONTH AGO", "3 MONTHS AGO"])
    n = 0
    for row in tbl[1:]:
        if len(row) < 6:
            continue
        ticker = row[0].strip().upper()
        if not ticker:
            continue
        ws.append([
            datetime(file_date.year, file_date.month, file_date.day),
            ticker, _int(row[1]), _int(row[2]),
            _int(row[3]), _int(row[4]), _int(row[5]),
        ])
        n += 1
    wb.save(out)
    return n


def generate_from_email(email, file_date: date) -> dict:
    """Write P5 and PTn xlsx files from an already-fetched Email object.

    Called by dispatch() during normal PS processing — no IMAP needed.
    Skips a file if it already exists in Archive (same precedence as PS).
    Returns a dict with keys 'p5' and 'ptn' → file path or 'skipped'.
    """
    tables = _parse_tables(email.html)
    p5_tbl  = next((t for t in tables if t and _is_p5(t[0])),  None)
    ptn_tbl = next((t for t in tables if t and _is_ptn(t[0])), None)
    ds = file_date.strftime("%Y-%m-%d")
    result: dict = {}

    if p5_tbl is None:
        log.warning("ps_grids: P5 table not found in %s", email.message_id)
        result["p5"] = "not found"
    else:
        out = P5_DIR / f"P5 {ds}.xlsx"
        if out.exists():
            result["p5"] = "skipped"
        else:
            n = _write_p5(p5_tbl, file_date, out)
            log.info("ps_grids: wrote %s (%d rows)", out, n)
            result["p5"] = str(out)

    if ptn_tbl is None:
        log.warning("ps_grids: PTn table not found in %s", email.message_id)
        result["ptn"] = "not found"
    else:
        out = PTN_DIR / f"PTn {ds}.xlsx"
        if out.exists():
            result["ptn"] = "skipped"
        else:
            n = _write_ptn(ptn_tbl, file_date, out)
            log.info("ps_grids: wrote %s (%d rows)", out, n)
            result["ptn"] = str(out)

    return result


def generate(target_date: date) -> None:
    ds = target_date.strftime("%Y-%m-%d")
    print(f"Connecting to Gmail IMAP...")
    mail = _connect()
    print(f"Searching for PS email on {ds}...")
    result = _fetch_ps_html(mail, target_date)
    mail.logout()

    if not result:
        print(f"ERROR: No PS email found for {ds}")
        return
    subj, html = result
    print(f"Found: {subj[:80]}")

    tables = _parse_tables(html)
    print(f"HTML tables in email: {len(tables)}")

    p5_tbl  = next((t for t in tables if t and _is_p5(t[0])),  None)
    ptn_tbl = next((t for t in tables if t and _is_ptn(t[0])), None)

    if p5_tbl is None:
        print("ERROR: P5 table not found")
    else:
        out = P5_DIR / f"P5 {ds}.xlsx"
        n = _write_p5(p5_tbl, target_date, out)
        print(f"Wrote {out} ({n} rows)")

    if ptn_tbl is None:
        print("ERROR: PTn table not found")
    else:
        out = PTN_DIR / f"PTn {ds}.xlsx"
        n = _write_ptn(ptn_tbl, target_date, out)
        print(f"Wrote {out} ({n} rows)")


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: python -m etl.hedgeye.ps_grids YYYY-MM-DD")
        sys.exit(1)
    try:
        d = date.fromisoformat(sys.argv[1])
    except ValueError:
        print(f"Invalid date: {sys.argv[1]}")
        sys.exit(1)
    generate(d)


if __name__ == "__main__":
    main()
